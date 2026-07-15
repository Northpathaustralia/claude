#!/usr/bin/env python3
"""Builds KAT_LASH_LOCK_ANNIVERSARY_PRESENTATION.pptx from 09_ANNIVERSARY_PRESENTATION.md.
Content is sourced verbatim from the markdown; this script only handles layout/styling.
"""
import re
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "09_ANNIVERSARY_PRESENTATION.md"
OUT = ROOT / "build" / "KAT_LASH_LOCK_ANNIVERSARY_PRESENTATION.pptx"

# Brand palette (04_BRAND_GUIDE.md)
PORCELAIN = RGBColor(0xFA, 0xF5, 0xF0)
ESPRESSO = RGBColor(0x2B, 0x1E, 0x1C)
BLUSH = RGBColor(0xF2, 0xC9, 0xC2)
GOLD = RGBColor(0xC9, 0xA2, 0x27)
SAGE = RGBColor(0xAE, 0xBB, 0xA8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

HEAD_FONT = "Fraunces"
BODY_FONT = "Inter"

STATEMENT_SLIDES = {1, 27, 28}  # full-bleed / centred statement layouts


def strip_md(text):
    """Strip light markdown emphasis markers, keep the words."""
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    text = re.sub(r"\*(.+?)\*", r"\1", text)
    text = re.sub(r"`(.+?)`", r"\1", text)
    return text.strip()


def parse_slides(md_text):
    blocks = re.split(r"\n\*\*SLIDE (\d+) — (.+?)\*\*\n", md_text)
    # blocks[0] is preamble; then triples of (num, heading, body) repeating
    slides = []
    for i in range(1, len(blocks), 3):
        num = int(blocks[i])
        heading = blocks[i + 1].strip()
        body_block = blocks[i + 2]
        fields = {}
        for key in ["Title", "Subtitle", "Body", "Speaker notes", "Visual", "Layout"]:
            m = re.search(rf"- \*\*{key}:\*\* (.+?)(?=\n- \*\*|\n\n|\Z)", body_block, re.S)
            fields[key] = strip_md(m.group(1).strip()) if m else ""
        slides.append({"num": num, "heading": heading, **fields})
    return slides


def set_background(slide, color):
    bg = slide.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = color


def add_gold_mark(slide, sw, sh, x=None, y=None):
    """Small arc-and-dot brand mark: a thin gold bar + dot, echoing the icon direction."""
    if x is None:
        x = sw - Inches(1.15)
    if y is None:
        y = sh - Inches(0.55)
    bar = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(0.55), Inches(0.06))
    bar.fill.solid()
    bar.fill.fore_color.rgb = GOLD
    bar.line.fill.background()
    bar.shadow.inherit = False
    dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, x + Inches(0.62), y - Inches(0.05), Inches(0.16), Inches(0.16))
    dot.fill.solid()
    dot.fill.fore_color.rgb = GOLD
    dot.line.fill.background()
    dot.shadow.inherit = False


def add_textbox(slide, x, y, w, h, text, size, color, bold=False, italic=False,
                 font=BODY_FONT, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, line_spacing=1.15):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = line_spacing
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.name = font
    run.font.color.rgb = color
    return box


def add_card(slide, x, y, w, h, fill=BLUSH, line_color=None):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    try:
        card.adjustments[0] = 0.045
    except Exception:
        pass
    card.fill.solid()
    card.fill.fore_color.rgb = fill
    if line_color:
        card.line.color.rgb = line_color
        card.line.width = Pt(0.75)
    else:
        card.line.fill.background()
    card.shadow.inherit = False
    return card


def build_standard_slide(slide, s, sw, sh):
    set_background(slide, PORCELAIN)
    margin = Inches(0.65)
    content_w = sw - 2 * margin

    # slide number chip
    add_textbox(slide, sw - Inches(1.0), Inches(0.3), Inches(0.6), Inches(0.3),
                f"{s['num']:02d}/28", 10, ESPRESSO, font=BODY_FONT, align=PP_ALIGN.RIGHT)

    # title (sized to safely wrap to 2 lines without colliding with the subtitle below)
    add_textbox(slide, margin, Inches(0.5), content_w * 0.62, Inches(1.5),
                s["Title"], 29, ESPRESSO, bold=True, font=HEAD_FONT, line_spacing=1.05)

    # subtitle chip
    if s["Subtitle"]:
        chip_w = min(Inches(6.4), Inches(0.35) * (len(s["Subtitle"]) * 0.09) + Inches(1.0))
        chip = add_card(slide, margin, Inches(2.05), content_w * 0.62, Inches(0.5), fill=BLUSH)
        tf = chip.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.18)
        tf.margin_right = Inches(0.18)
        tf.margin_top = Inches(0.02)
        tf.margin_bottom = Inches(0.02)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        run = p.add_run()
        run.text = s["Subtitle"]
        run.font.size = Pt(14)
        run.font.italic = True
        run.font.name = BODY_FONT
        run.font.color.rgb = ESPRESSO

    # body
    if s["Body"] and s["Body"] != "(none — let it breathe)":
        add_textbox(slide, margin, Inches(2.85), content_w * 0.62, Inches(3.5),
                    s["Body"], 15.5, ESPRESSO, font=BODY_FONT, line_spacing=1.3)

    # visual placeholder card (right column)
    vx = margin + content_w * 0.66
    vw = content_w * 0.34
    vcard = add_card(slide, vx, Inches(1.55), vw, Inches(4.35), fill=RGBColor(0xF3, 0xEC, 0xE6), line_color=BLUSH)
    tf = vcard.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.22)
    tf.margin_right = Inches(0.22)
    tf.margin_top = Inches(0.22)
    tf.margin_bottom = Inches(0.22)
    tf.vertical_anchor = MSO_ANCHOR.TOP
    p0 = tf.paragraphs[0]
    r0 = p0.add_run()
    r0.text = "VISUAL DIRECTION"
    r0.font.size = Pt(10)
    r0.font.bold = True
    r0.font.name = BODY_FONT
    r0.font.color.rgb = GOLD
    p1 = tf.add_paragraph()
    p1.space_before = Pt(6)
    r1 = p1.add_run()
    r1.text = s["Visual"]
    r1.font.size = Pt(12.5)
    r1.font.italic = True
    r1.font.name = BODY_FONT
    r1.font.color.rgb = ESPRESSO
    p1.line_spacing = 1.25
    if s["Layout"]:
        p2 = tf.add_paragraph()
        p2.space_before = Pt(14)
        r2 = p2.add_run()
        r2.text = f"Layout: {s['Layout']}"
        r2.font.size = Pt(10)
        r2.font.name = BODY_FONT
        r2.font.color.rgb = RGBColor(0x7A, 0x6A, 0x64)

    add_gold_mark(slide, sw, sh)


def build_statement_slide(slide, s, sw, sh):
    set_background(slide, ESPRESSO)
    margin = Inches(1.1)
    content_w = sw - 2 * margin

    add_textbox(slide, sw - Inches(1.0), Inches(0.3), Inches(0.6), Inches(0.3),
                f"{s['num']:02d}/28", 10, RGBColor(0xC9, 0xB8, 0xB0), font=BODY_FONT, align=PP_ALIGN.RIGHT)

    add_textbox(slide, margin, Inches(2.15), content_w, Inches(1.7),
                s["Title"], 40, PORCELAIN, bold=True, font=HEAD_FONT,
                align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.05)

    if s["Subtitle"]:
        add_textbox(slide, margin, Inches(3.75), content_w, Inches(0.6),
                    s["Subtitle"], 16, BLUSH, italic=True, font=BODY_FONT,
                    align=PP_ALIGN.CENTER)

    if s["Body"] and s["Body"] != "(none — let it breathe)":
        add_textbox(slide, margin + content_w * 0.12, Inches(4.55), content_w * 0.76, Inches(2.0),
                    s["Body"], 15, RGBColor(0xE9, 0xDE, 0xD8), font=BODY_FONT,
                    align=PP_ALIGN.CENTER, line_spacing=1.3)

    # gold arc + dot centred above title
    bar_w = Inches(0.9)
    bar = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, sw / 2 - bar_w / 2, Inches(1.75), bar_w, Inches(0.05))
    bar.fill.solid()
    bar.fill.fore_color.rgb = GOLD
    bar.line.fill.background()
    bar.shadow.inherit = False
    dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, sw / 2 - Inches(0.09), Inches(1.55), Inches(0.18), Inches(0.18))
    dot.fill.solid()
    dot.fill.fore_color.rgb = GOLD
    dot.line.fill.background()
    dot.shadow.inherit = False


def main():
    md_text = SRC.read_text(encoding="utf-8")
    slides_data = parse_slides(md_text)
    assert len(slides_data) == 28, f"expected 28 slides, parsed {len(slides_data)}"

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    for s in slides_data:
        slide = prs.slides.add_slide(blank_layout)
        if s["num"] in STATEMENT_SLIDES:
            build_statement_slide(slide, s, prs.slide_width, prs.slide_height)
        else:
            build_standard_slide(slide, s, prs.slide_width, prs.slide_height)

        notes = slide.notes_slide
        notes.notes_text_frame.text = s["Speaker notes"] or ""

    OUT.parent.mkdir(exist_ok=True)
    prs.save(OUT)
    print(f"Wrote {OUT} ({len(slides_data)} slides)")


if __name__ == "__main__":
    main()
