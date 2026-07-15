#!/usr/bin/env python3
"""Builds LETTER_TO_KAT.docx from 10_FOUNDER_LETTER_TO_KAT.md.
Content is sourced verbatim; this script only handles layout/styling.
No brand colours or logos — a plain, warm letter, per 11_CODEX_BUILD_INSTRUCTIONS.md.
"""
import re
from pathlib import Path

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "10_FOUNDER_LETTER_TO_KAT.md"
OUT = ROOT / "build" / "LETTER_TO_KAT.docx"

INK = RGBColor(0x2A, 0x25, 0x22)
BODY_FONT = "Garamond"


def add_run(p, text, italic=False, bold=False, size=13):
    run = p.add_run(text)
    run.italic = italic
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = BODY_FONT
    run.font.color.rgb = INK
    return run


def render_inline(p, text, size=13):
    """Render **bold** spans within a paragraph; everything else plain."""
    parts = re.split(r"(\*\*.+?\*\*)", text)
    for part in parts:
        if part.startswith("**") and part.endswith("**"):
            add_run(p, part[2:-2], bold=True, size=size)
        else:
            add_run(p, part, size=size)


def main():
    lines = SRC.read_text(encoding="utf-8").splitlines()

    # Strip the H1 title and the italic instruction line at top; body starts after the --- rule
    body_start = 0
    for i, line in enumerate(lines):
        if line.strip() == "---":
            body_start = i + 1
            break
    paragraphs = [l for l in lines[body_start:] if l.strip() != ""]

    doc = Document()
    section = doc.sections[0]
    section.page_height = Inches(11.69)  # A4
    section.page_width = Inches(8.27)
    section.top_margin = Inches(1.3)
    section.bottom_margin = Inches(1.3)
    section.left_margin = Inches(1.4)
    section.right_margin = Inches(1.4)

    style = doc.styles["Normal"]
    style.font.name = BODY_FONT
    style.font.size = Pt(13)
    style.font.color.rgb = INK
    style.paragraph_format.space_after = Pt(14)
    style.paragraph_format.line_spacing = 1.35

    for raw in paragraphs:
        text = raw.strip()
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT

        if text.startswith("*[") and text.endswith("]*"):
            render_inline(p, text[1:-1], size=12)
            p.runs[-1].italic = True
            for r in p.runs:
                r.italic = True
                r.font.color.rgb = RGBColor(0x8A, 0x7E, 0x78)
            continue

        if text == "Kat,":
            render_inline(p, text, size=16)
            p.runs[0].bold = False
            continue

        if text in ("James",):
            render_inline(p, text, size=14)
            continue

        render_inline(p, text)

    OUT.parent.mkdir(exist_ok=True)
    doc.save(OUT)
    print(f"Wrote {OUT} ({len(paragraphs)} paragraphs)")


if __name__ == "__main__":
    main()
